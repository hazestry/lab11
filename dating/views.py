from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from django.db.models import Q, Count
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

from .models import UserProfile, Message, Like
from .forms import UserProfileForm, MessageForm, ProfileSearchForm
from users.models import CustomUser


def home_view(request):
    """
    Главная страница
    """
    return render(request, 'dating/home.html')


@login_required
def profile_list_view(request):
    profiles = UserProfile.objects.filter(is_visible=True).select_related('user')
    form = ProfileSearchForm(request.GET)

    if form.is_valid():
        if form.cleaned_data.get('name'):
            name = form.cleaned_data['name']
            profiles = profiles.filter(
                Q(user__first_name__icontains=name) |
                Q(user__last_name__icontains=name)
            )

        if form.cleaned_data.get('age_min'):
            profiles = profiles.filter(age__gte=form.cleaned_data['age_min'])

        if form.cleaned_data.get('age_max'):
            profiles = profiles.filter(age__lte=form.cleaned_data['age_max'])

        if form.cleaned_data.get('city'):
            profiles = profiles.filter(city__icontains=form.cleaned_data['city'])

        if form.cleaned_data.get('education'):
            profiles = profiles.filter(education=form.cleaned_data['education'])

        if form.cleaned_data.get('relationship_goal'):
            profiles = profiles.filter(relationship_goal=form.cleaned_data['relationship_goal'])

    if hasattr(request.user, 'dating_profile'):
        profiles = profiles.exclude(user=request.user)

    context = {
        'profiles': profiles,
        'form': form,
        'is_guest': request.user.is_guest() if hasattr(request.user, 'is_guest') else False,
    }

    return render(request, 'dating/profile_list.html', context)


@login_required
def profile_detail_view(request, pk):
    profile = get_object_or_404(UserProfile, pk=pk)

    user_liked = False
    if not request.user.is_guest():
        user_liked = Like.objects.filter(
            user_from=request.user,
            user_to=profile.user
        ).exists()

    context = {
        'profile': profile,
        'user_liked': user_liked,
        'is_guest': request.user.is_guest() if hasattr(request.user, 'is_guest') else False,
    }

    return render(request, 'dating/profile_detail.html', context)


@login_required
def create_profile_view(request):
    # У пользователя уже есть профиль
    if hasattr(request.user, 'dating_profile'):
        messages.info(request, 'У вас уже есть профиль')
        return redirect('dating:edit_profile')

    if request.user.is_guest():
        messages.error(request, 'Гости не могут создавать профили. Зарегистрируйтесь!')
        return redirect('users:register')

    if request.method == 'POST':
        form = UserProfileForm(request.POST)
        if form.is_valid():
            profile = form.save(commit=False)
            profile.user = request.user
            profile.save()
            messages.success(request, 'Профиль успешно создан!')
            return redirect('dating:profile_list')
    else:
        form = UserProfileForm()

    return render(request, 'dating/profile_form.html', {'form': form, 'title': 'Создать профиль'})


@login_required
def edit_profile_view(request):
    if request.user.is_guest():
        messages.error(request, 'Гости не могут редактировать профили')
        return redirect('dating:profile_list')

    profile = get_object_or_404(UserProfile, user=request.user)

    if request.method == 'POST':
        form = UserProfileForm(request.POST, instance=profile)
        if form.is_valid():
            form.save()
            messages.success(request, 'Профиль успешно обновлен!')
            return redirect('dating:my_profile')
    else:
        form = UserProfileForm(instance=profile)

    return render(request, 'dating/profile_form.html', {'form': form, 'title': 'Редактировать профиль'})


@login_required
def my_profile_view(request):
    if not hasattr(request.user, 'dating_profile'):
        messages.info(request, 'Создайте свой профиль')
        return redirect('dating:create_profile')

    profile = request.user.dating_profile

    likes_received = Like.objects.filter(user_to=request.user).count()
    likes_given = Like.objects.filter(user_from=request.user).count()
    messages_count = Message.objects.filter(
        Q(sender=request.user) | Q(recipient=request.user)
    ).count()

    context = {
        'profile': profile,
        'likes_received': likes_received,
        'likes_given': likes_given,
        'messages_count': messages_count,
    }

    return render(request, 'dating/my_profile.html', context)


@login_required
@require_http_methods(["POST"])
def toggle_like_view(request, user_id):
    if request.user.is_guest():
        return JsonResponse({'error': 'Гости не могут ставить лайки'}, status=403)

    target_user = get_object_or_404(CustomUser, id=user_id)

    if target_user == request.user:
        return JsonResponse({'error': 'Нельзя лайкнуть самого себя'}, status=400)

    like, created = Like.objects.get_or_create(
        user_from=request.user,
        user_to=target_user
    )

    if not created:
        like.delete()
        return JsonResponse({'status': 'unliked', 'message': 'Лайк убран'})

    return JsonResponse({'status': 'liked', 'message': 'Лайк поставлен'})


@login_required
def messages_view(request, user_id=None):
    if request.user.is_guest():
        messages.error(request, 'Гости не могут отправлять сообщения. Зарегистрируйтесь!')
        return redirect('users:register')

    conversations = Message.objects.filter(
        Q(sender=request.user) | Q(recipient=request.user)
    ).values('sender', 'recipient').distinct()

    conversation_users_ids = set()
    for conv in conversations:
        if conv['sender'] != request.user.id:
            conversation_users_ids.add(conv['sender'])
        if conv['recipient'] != request.user.id:
            conversation_users_ids.add(conv['recipient'])

    conversation_users = CustomUser.objects.filter(id__in=conversation_users_ids)

    current_chat_user = None
    chat_messages = []
    form = None

    if user_id:
        current_chat_user = get_object_or_404(CustomUser, id=user_id)

        chat_messages = Message.objects.filter(
            (Q(sender=request.user) & Q(recipient=current_chat_user)) |
            (Q(sender=current_chat_user) & Q(recipient=request.user))
        ).order_by('created_at')

        Message.objects.filter(
            sender=current_chat_user,
            recipient=request.user,
            is_read=False
        ).update(is_read=True)

        if request.method == 'POST':
            form = MessageForm(request.POST)
            if form.is_valid():
                message = form.save(commit=False)
                message.sender = request.user
                message.recipient = current_chat_user
                message.save()
                messages.success(request, 'Сообщение отправлено!')
                return redirect('dating:messages', user_id=user_id)
        else:
            form = MessageForm()

    context = {
        'conversation_users': conversation_users,
        'current_chat_user': current_chat_user,
        'chat_messages': chat_messages,
        'form': form,
    }

    return render(request, 'dating/messages.html', context)


@staff_member_required
def export_data_view(request):

    if request.method == 'POST':
        export_profiles = request.POST.get('export_profiles') == 'on'
        export_messages = request.POST.get('export_messages') == 'on'
        export_likes = request.POST.get('export_likes') == 'on'

        wb = Workbook()
        wb.remove(wb.active)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # --- ПРОФИЛИ ---
        if export_profiles:
            ws_profiles = wb.create_sheet("Профили")
            headers = ['ID', 'Пользователь', 'Возраст', 'Город', 'Образование',
                       'Профессия', 'Интересы', 'Цель', 'Статус', 'Дата создания', 'Дата обновления']
            ws_profiles.append(headers)

            for cell in ws_profiles[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            for profile in UserProfile.objects.all().select_related('user'):
                ws_profiles.append([
                    profile.id,
                    profile.user.get_full_name(),
                    profile.age,
                    profile.city,
                    profile.education,
                    profile.profession,
                    profile.interests,
                    profile.relationship_goal,
                    profile.status,
                    profile.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                    profile.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
                ])

            for column in ws_profiles.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                ws_profiles.column_dimensions[column_letter].width = min(max_length + 2, 50)

        # --- СООБЩЕНИЯ ---
        if export_messages:
            ws_messages = wb.create_sheet("Сообщения")
            headers = ['ID', 'Отправитель', 'Получатель', 'Текст', 'Прочитано', 'Дата отправки']
            ws_messages.append(headers)

            for cell in ws_messages[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            for message in Message.objects.all().select_related('sender', 'recipient'):
                ws_messages.append([
                    message.id,
                    message.sender.get_full_name(),
                    message.recipient.get_full_name(),
                    message.text[:100],
                    'Да' if message.is_read else 'Нет',
                    message.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                ])

            for column in ws_messages.columns:
                max_length = 0
                col = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                ws_messages.column_dimensions[col].width = min(max_length + 2, 50)

        # --- ЛАЙКИ ---
        if export_likes:
            ws_likes = wb.create_sheet("Лайки")
            headers = ['ID', 'От пользователя', 'Кому', 'Дата']
            ws_likes.append(headers)

            for cell in ws_likes[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            for like in Like.objects.all().select_related('user_from', 'user_to'):
                ws_likes.append([
                    like.id,
                    like.user_from.get_full_name(),
                    like.user_to.get_full_name(),
                    like.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                ])

            for column in ws_likes.columns:
                max_length = 0
                col = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                ws_likes.column_dimensions[col].width = min(max_length + 2, 50)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'swipeheart_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response['Content-Disposition'] = f'attachment; filename={filename}'

        wb.save(response)
        return response

    stats = {
        'profiles_count': UserProfile.objects.count(),
        'messages_count': Message.objects.count(),
        'likes_count': Like.objects.count(),
    }

    return render(request, 'dating/export_data.html', {'stats': stats})
